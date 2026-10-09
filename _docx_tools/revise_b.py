"""批B/批C：AI 专项方案逻辑一致性修复。

批B：文本级逻辑修复 + 映射表补编号（追溯断链缝合）
批C：结构性插入（含 1.5 建设范围边界——落地『AI 不做对账』决策）

所有操作带期望命中数校验，任一不符则中止且不落盘。
"""
import sys
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph

DOC = "../AI专项方案v2.docx"


def all_runs(doc):
    for p in doc.paragraphs:
        for r in p.runs:
            yield p, r
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        yield p, r


def rep(doc, old, new, expect):
    hits = 0
    for _, r in all_runs(doc):
        if old in r.text:
            r.text = r.text.replace(old, new)
            hits += 1
    flag = "✓" if hits == expect else "✗"
    print(f"  {flag} rep {hits}/{expect} {old[:34]!r}")
    return hits == expect


def cell(doc, ti, ri, ci, old, new):
    c = doc.tables[ti].rows[ri].cells[ci]
    for p in c.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                print(f"  ✓ cell T{ti}r{ri}c{ci} {new[:40]!r}")
                return True
    print(f"  ✗ cell T{ti}r{ri}c{ci} 未命中 {old!r}")
    return False


def find(doc, sub):
    for p in doc.paragraphs:
        if sub in p.text:
            return p
    return None


def clone_after(anchor, like, bold_txt, norm_txt, before=False):
    """复制 like 段落的段落样式与 run 格式，插入到 anchor 之后（或之前）。

    必须剔除副本里的书签与批注标记：Heading 段落带 _TocXXXXX 书签，
    直接 deepcopy 会造重复书签 ID，破坏 Word 目录域的跳转。
    """
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    el = deepcopy(like._p)
    for tag in ("r", "bookmarkStart", "bookmarkEnd",
                "commentRangeStart", "commentRangeEnd"):
        nodes = list(el.findall(W + tag))
        if tag == "r":
            nodes = nodes[2:]          # 只保留前两个 run 作为格式载体
        for n in nodes:
            el.remove(n)
    if before:
        anchor._p.addprevious(el)
    else:
        anchor._p.addnext(el)
    np = Paragraph(el, like._parent)
    if norm_txt is None:
        np.runs[0].text = bold_txt
    else:
        if len(np.runs) == 1:
            np.runs[0].text = bold_txt + norm_txt
        else:
            np.runs[0].text = bold_txt
            np.runs[1].text = norm_txt
            for r in np.runs[2:]:
                r.text = ""
    return np


def insert_block(doc, anchor_sub, like_sub, items, before=False):
    anchor, like = find(doc, anchor_sub), find(doc, like_sub)
    if anchor is None or like is None:
        print(f"  ✗ 锚点缺失: {anchor_sub!r} / {like_sub!r}")
        return False
    cur = anchor
    for b, n in items:
        cur = clone_after(cur, like, b, n, before=before)
    print(f"  ✓ 插入 {len(items)} 段（锚点 {anchor_sub[:22]!r}）")
    return True


# ───────────────────────── 批B ─────────────────────────
def batch_b(doc):
    ok = True
    # B2 数据闭环措辞与「AI 不替业务决策」边界对齐
    ok &= rep(doc, "→ 商品中台的供应商推品策略优化依据",
              "→ 经公共底座数据中台分析后形成的供应商推品策略优化依据（供运营决策，非 AI 自动执行）", 1)
    ok &= rep(doc, "→ 商品中台的商品池补品建议",
              "→ 经公共底座数据中台分析后形成的商品池补品建议（供选品决策参考）", 1)
    ok &= rep(doc, "→ 商品中台的运营决策依据",
              "→ 商品中台的运营决策依据（由数据中台分析产出，供运营决策）", 1)
    ok &= rep(doc, "反向指导供应商推品策略，淘汰长期低通过率的供应商商品。",
              "反向指导供应商推品策略，对长期低通过率的供应商商品提出淘汰建议（分析产出由数据中台承载，处置决策由业务方做出）。", 1)
    # B3 客服调用顺序与「规则优先、成本可控」原则自洽
    ok &= rep(doc, "用户咨询进入后先做意图识别",
              "用户咨询进入后先执行转人工与敏感词硬规则（规则引擎前置，命中即转人工并留痕，不产生大模型调用成本）；未命中硬规则时再做意图识别", 1)
    ok &= rep(doc, "任何场景识别到情感异常或转人工关键词即转人工。",
              "会话过程中持续监控情感异常与转人工关键词，命中即转人工。", 1)
    # B6 无基线的效果承诺补口径
    ok &= rep(doc, "将人工上架工作量降低 60% 以上",
              "将人工上架工作量显著下降（目标降幅 60% 以上，〔示意目标，POC 阶段锁定人工处理基线后转为正式验收值〕）", 1)
    # B7 标注方法学定义，消除「自动化标注评测集」与 8.2 交叉验证原则的冲突
    ok &= rep(doc, "评测集标注由专职标注团队完成，多人交叉验证，标注一致性 ≥ 90%",
              "评测集标注采用“ML 预标注 + 人工交叉复核”的半自动方式，由专职标注团队执行并双人交叉验证，标注一致性 ≥ 90%；"
              "自动预标注只用于提升效率，最终判定必须由人工复核做出——禁止用模型自行标注的结果评价同一模型，否则训测同源导致指标虚高", 1)
    # B11 按 1.5 边界收口，长期演进不再留对账回潮的口子
    ok &= rep(doc, "：从文本扩展到图像（商品图片识别、发票 OCR）、语音（客服语音转文字、语音搜索）",
              "：从文本扩展到图像（商品图片识别、包装与资质证照核验）与语音（客服语音转文字、语音搜索）；"
              "单据图像类能力仅作数据录入辅助，不介入对账与金额判定（见 1.5 建设范围边界）", 1)
    # B9 缝合「技术层与指标层有比价、功能清单无比价」的断链
    ok &= rep(doc, "AIP-015 至 AIP-031 功能清单。",
              "AIP-015 至 AIP-031 功能清单。其中同款比价不单列功能编号，由 AIP-010 同款精确匹配（产出同款簇）"
              "与 AIP-020 属性聚合筛选（聚合多供应商报价）联合承载，技术实现见 5.3.3 比价聚合模型，效果度量见 5.3.4 比价结果采纳率。", 1)

    # B10 7.1 映射表补功能编号，打通「业务功能 → 编号」双向追溯
    t5 = [
        (1, 0, "商品库管理", "商品库管理（AIP-001/002）"),
        (2, 0, "商品类目治理", "商品类目治理（AIP-003/004）"),
        (3, 0, "商品属性治理", "商品属性治理（AIP-005/006）"),
        (4, 0, "商品智能体-审核", "商品智能体-审核（AIP-007/009）"),
        (5, 0, "商品智能体-规范性校验", "商品智能体-规范性校验（AIP-008）"),
        (6, 0, "商品智能体-同款匹配", "商品智能体-同款匹配（AIP-010）"),
        (7, 0, "商品智能体-SEO 标题", "商品智能体-SEO 标题（AIP-011）"),
        (8, 0, "商品智能体-属性提取", "商品智能体-属性提取（AIP-012）"),
        (9, 0, "商品智能体-合理性检测", "商品智能体-合理性检测（AIP-013）"),
        (10, 0, "商品智能体-集采合规", "商品智能体-集采合规（AIP-014）"),
        (11, 0, "商品搜索", "商品搜索（AIP-015 至 AIP-023）"),
        (12, 0, "对话导购", "对话导购（AIP-024 至 AIP-028）"),
        (13, 0, "猜你喜欢", "猜你喜欢（AIP-029 至 AIP-031）"),
        (14, 0, "在线客服", "在线客服（AICS-001 至 AICS-009）"),
        (15, 0, "人工接待", "人工接待（AICS-010 至 AICS-015）"),
        (16, 0, "智能体接待", "智能体接待（AICS-016 至 AICS-025）"),
        (17, 0, "报表分析", "报表分析（AICS-029；基础设置 AICS-026 至 AICS-028）"),
    ]
    for ri, ci, old, new in t5:
        ok &= cell(doc, 5, ri, ci, old, new)
    # B8-前置：8.5 评测节奏责任方与 9.1 组织角色对齐（角色补齐在批C）
    ok &= cell(doc, 6, 1, 2, "算法团队", "AI 算法/工程负责人 + 外部 AI 服务商")
    ok &= cell(doc, 6, 2, 2, "算法团队 +  产品团队", "AI 算法/工程负责人 + 产品团队")
    return ok


# ───────────────────────── 批C ─────────────────────────
def batch_c(doc):
    ok = True
    # C1 新增 1.5 建设范围边界（落地用户决策：AI 不做任何对账能力）
    #    必须插在「二、AI 中台总体架构」之前，且标题用 Heading 2 以进目录域
    ok &= insert_block(doc, "二、AI 中台总体架构", "1.4 技术路线总览", [
        ("1.5 建设范围边界（本专项不做什么）", None),
    ], before=True)
    h = find(doc, "1.5 建设范围边界")
    if h is None:
        print("  ✗ C1 标题未生成")
        return False
    clone_after(h, find(doc, "本方案是商城整体咨询项目中的 AI 子方案"),
                "AI 能力的边界与能力本身同等重要。为避免范围蔓延与投入错配，本专项明确以下三条不做，"
                "并要求方案内外所有相关表述与本节保持一致：", None)
    lead = find(doc, "AI 能力的边界与能力本身同等重要")
    like = find(doc, "上架效率提升")
    items = [
        ("确定性财务判定不做 AI。",
         "对账、三单匹配（订单—入库—发票）、金额与数量的校验属于确定性计算与强审计域：规则引擎可 100% 覆盖判定逻辑，"
         "大模型介入只会引入幻觉风险与资损敞口，逐单调用更与“按次付费”的成本约束直接冲突。"
         "本专项不建设 AI 对账能力，单据 OCR 亦不纳入本期；对账效率问题由商城系统的单据直连、结构化校验与规则引擎改造解决，"
         "属于商城系统升级范围而非 AI 专项范围。"),
        ("新业务场景的系统形态改造不做 AI。",
         "员工福利、对外销售等场景的“系统不适配”是流程与功能缺失，由商城产品功能建设解决；"
         "三大 AI 中台为其提供可复用的商品治理、搜索导购与客服能力，不为其单独建设 AI 功能点。"),
        ("经营决策类分析不做独立 AI 功能。",
         "零结果 Query 的补品建议、低通过率供应商的淘汰建议等属于数据分析产出，"
         "由公共底座的数据中台与报表分析能力承载（见 6.2.4 与 7.1），作为运营输入供业务方决策，不以 AI 功能点形式交付。"),
        ("一句话原则。",
         "AI 只承担规则无法 100% 覆盖的语义理解与判断类工作；凡能被规则或结构化计算确定的，一律不使用 AI。"
         "该原则同时约束技术选型、成本口径与本方案后续章节的能力清单。"),
    ]
    cur = lead
    for b, n in items:
        cur = clone_after(cur, like, b, n)
    print("  ✓ C1 1.5 范围边界小节已插入" if ok else "  ✗ C1")

    # C2 公共底座职责边界（消除「避免重复建设」与各章自建规则/评测集的表观矛盾）
    anchor = find(doc, "评测平台：统一的评测集管理")
    clone_after(anchor, find(doc, "本方案是商城整体咨询项目中的 AI 子方案"),
                "职责边界：公共底座提供的是引擎与平台（规则执行引擎、模型推理服务、评测计算框架、数据管道），"
                "三大中台维护的是内容资产（规则集、模型与提示词、评测集与指标定义）。"
                "底座不替代中台的业务规则与评测标准，中台不重复建设底层能力——这是“中台”相对于“功能”的成本分界。", None)
    print("  ✓ C2 底座职责边界")

    # C3 数据与基础设施层补部署口径（对齐 ADR-014：前期纯云端 + 脱敏网关，存储全内网）
    a = find(doc, "向量检索引擎：支持商品向量、对话向量的高效检索")
    clone_after(a, find(doc, "本方案是商城整体咨询项目中的 AI 子方案"),
                "部署口径：本期不建设本地大模型与 GPU 推理集群，大模型能力统一经云端 API 网关接入并支持多供应商主备切换；"
                "商品、价格、订单、对话等原始数据的存储与处理全部保留在内网，送云端的内容经脱敏与字段最小化处理。"
                "“数据不出域”是本专项的硬约束，约束的是数据流转边界而非算力位置。", None)
    print("  ✓ C3 部署口径")

    # C4 9.1 组织架构补齐技术与标注角色（与 8.5 责任方、9.4 供应商边界闭环）
    ok &= insert_block(doc, "运营人员（兼职，1-2人）", "AI 产品负责人", [
        ("AI 算法/工程负责人（1 人）：",
         "负责提示词与模型实现、评测集构建与离线评测、灰度发布与回滚，并对外部 AI 服务商的交付质量做技术验收。"
         "该角色是 8.5 评测节奏中离线评测与 A/B 实验的责任主体，不可由业务或运营岗位兼任。"),
        ("数据标注（外部供应商承担，内部验收）：",
         "评测集与训练集的大规模标注工作量由外部数据标注供应商承担（见 9.4），"
         "由 AI 算法/工程负责人按 8.2 的交叉验证标准验收标注质量，标注规范与判定口径归内部所有。"),
    ])
    print("  ✓ C4 组织角色补齐")

    # C5 8.2 增补评测集版本冻结原则（堵住 bad case 回流训练集后跨期指标不可比的风险）
    ok &= insert_block(doc, "持续更新：评测集定期补充新场景、新 bad case", "持续更新：评测集定期补充新场景、新 bad case", [
        ("版本冻结与换代：",
         "评测集按版本冻结管理，一个评测周期内不得增删样本。线上 bad case 回流训练集时，"
         "同步评估是否需要换代评测集；换代后新旧指标不直接对比，需并行评测一版建立换算关系，"
         "否则跨期指标不可比，会用“评测集变难”掩盖真实的效果回退。"),
    ])
    print("  ✓ C5 评测集版本冻结")
    return ok


def main():
    which = sys.argv[1] if len(sys.argv) > 1 else "b"
    doc = Document(DOC)
    ok = batch_b(doc) if which == "b" else batch_c(doc)
    if not ok:
        print("ABORT: 存在未命中操作，未保存")
        sys.exit(1)
    doc.save(DOC)
    print("saved:", DOC, "| batch", which)


if __name__ == "__main__":
    main()
