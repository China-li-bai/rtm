"""批F：自审修正——批E 引入的计数矛盾 + 上轮漏扫的表格残留。

1) 1.5 引导句写「以下三条不做」，批E 增第四条后实际为四条 → 自相矛盾（正是本次批评文档的那类错误）
2) 1.3 的 T1 定位表「下游输出」列仍含「坐席辅助」——上轮只扫正文 paragraphs，漏了表格
"""
from docx import Document
from revise_b import rep, cell

DOC = "../AI专项方案v2.docx"

doc = Document(DOC)
ok = True

ok &= rep(doc, "本专项明确以下三条不做", "本专项明确以下四条不做", 1)

# T1 = 1.3 三大中台定位表；客服中台在第 4 行（r0 表头）的「下游输出」列
ok &= cell(doc, 1, 3, 3, "智能问答、工单分派、坐席辅助、转人工",
           "智能问答、工单分派、转人工（坐席辅助属远期方向，不在本期功能清单）")

if not ok:
    print("ABORT: 未保存")
    raise SystemExit(1)
doc.save(DOC)
print("saved:", DOC)

# 复测：全文（含表格）不应再有未降级的「坐席辅助」
txt = [p.text for p in doc.paragraphs]
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            txt.append(c.text)
bad = [x[:80] for x in "\n".join(txt).split("\n")
       if "坐席辅助" in x and "远期" not in x and "不作交付承诺" not in x]
print("未降级的坐席辅助行:", bad or "NONE")
print("引导句计数:", [x[:40] for x in txt if "本专项明确以下" in x])
