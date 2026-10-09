"""批D：模块边界修复（跨中台重复定义）+ 措辞含糊收口 + 目录域自动刷新。

1) 生鲜预占规则在 4.3.2（商品）与 5.3.2（搜索）被重复定义，且预占本质属交易/库存域
   → 商品侧保留唯一定义并标注共享归属，搜索侧降级为只读引用
2) 跨中台复用资产（词库/同义词库/属性图谱）归属数据中台统一版本管理
3) 「可比性」含糊措辞去模糊
4) 9.3 风险表补「范围蔓延」，与新增的 1.5 边界形成闭环
5) settings.updateFields=true：Word 打开时自动刷新目录域（1.5 新节 + 页码）
"""
from copy import deepcopy
from docx import Document

DOC = "../AI专项方案v2.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def find(doc, sub):
    for p in doc.paragraphs:
        if sub in p.text:
            return p
    return None


def rep(doc, old, new, expect=1):
    hits = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                hits += 1
    print(f"  {'✓' if hits == expect else '✗'} rep {hits}/{expect} {old[:32]!r}")
    return hits == expect


def rep_in(doc, para_sub, old, new):
    """只在首个匹配 para_sub 的段落内替换，避免同名串误伤其他章节。"""
    p = find(doc, para_sub)
    if p is None:
        print(f"  ✗ 段落锚点缺失 {para_sub[:30]!r}")
        return False
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            print(f"  ✓ rep_in [{para_sub[:18]!r}] {old[:26]!r}")
            return True
    print(f"  ✗ rep_in 未命中 {old[:30]!r} in {p.text[:40]!r}")
    return False


def add_risk_row(table, cols):
    """复制末行 tr 以保持边框/字体一致，再改写单元格文本。"""
    tr = deepcopy(table.rows[-1]._tr)
    table.rows[-1]._tr.addnext(tr)
    row = table.rows[-1]
    for cell, text in zip(row.cells, cols):
        ps = cell.paragraphs
        got = False
        for r in ps[0].runs:
            if not got:
                r.text = text
                got = True
            else:
                r.text = ""
        if not got and ps[0].runs == []:
            ps[0].add_run(text)
    print("  ✓ 风险表新增行:", cols[0])
    return True


def set_update_fields(doc):
    """按 ECMA-376 w:CT_Settings 顺序位插入 w:updateFields，避免 Word 拒文档。

    updateFields 在序列中位于 hdrShapeDefaults / footnotePr / endnotePr / compat 之前，
    插到开头或末尾均可能违反顺序，因此找第一个后继元素前插入。
    """
    st = doc.settings.element
    if st.find(W + "updateFields") is not None:
        print("  · updateFields 已存在")
        return True
    after = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "rsids",
             "docVars", "shapeDefaults", "stdDocVars", "themeFontLang", "clrSchemeMapping"]
    anchor = None
    for name in after:
        anchor = st.find(W + name)
        if anchor is not None:
            break
    uf = st.makeelement(W + "updateFields", {W + "val": "true"})
    if anchor is not None:
        anchor.addprevious(uf)
        print(f"  ✓ updateFields=true（插入于 {name} 前，符合 schema 顺序）")
    else:
        st.append(uf)
        print("  ✓ updateFields=true（追加于末尾）")
    return True


def main():
    doc = Document(DOC)
    ok = True

    # ── 1 生鲜预占规则：域归属 + 去重
    ok &= rep_in(doc, "生鲜类商品的预占有效期与其他品类不同",
                 "生鲜预占规则", "生鲜预占规则（跨域共享·唯一定义处）")
    ok &= rep_in(doc, "生鲜类商品的预占有效期与其他品类不同",
                 "，规则引擎按品类配置预占逻辑。",
                 "。预占本质是交易与库存域的硬性约束，不属于商品治理域：该规则由公共底座的规则引擎统一持有、配置与执行，"
                 "商品中台与搜索中台按品类只读引用，不在中台内重复定义（见 2.4 职责边界）。")
    ok &= rep_in(doc, "搜索生鲜品类时，前置展示预占时效提示",
                 "生鲜预占规则", "生鲜预占规则（引用共享规则）")
    ok &= rep_in(doc, "搜索生鲜品类时，前置展示预占时效提示",
                 "下单时按生鲜预占规则锁定库存。",
                 "预占与库存锁定由交易域按 4.3.2 的共享规则执行，搜索中台只做时效提示与参数透传，不重复定义该规则。")

    # ── 2 跨中台复用资产归属数据中台
    ok &= rep(doc, "工单库、评测集仓库等。",
              "工单库、评测集仓库等。其中跨中台复用的资产（违规违禁词库、违禁图片样本库、客服敏感词库、"
              "同义词库、品类属性图谱）在此统一版本管理，各中台按需引用，避免同一资产多处维护导致口径漂移。")

    # ── 3 「可比性」去模糊（注意 P24 的“构建可比性亮点”是独立加粗 run，只改该 run 文本以保留样式）
    ok &= rep(doc, "本专项方案聚焦三个最具落地价值且最具可比性的能力方向",
              "本专项方案聚焦三个落地价值最高、且效果可对外对标验证的能力方向")
    ok &= rep(doc, "构建可比性亮点", "打造可对标的差异化亮点")
    ok &= rep(doc, "形成可对内对外展示的案例。", "形成可对内汇报、可对外对标验证的标杆案例。")

    # ── 4 风险表补「范围蔓延」，与 1.5 边界闭环
    add_risk_row(doc.tables[7], (
        "范围蔓延",
        "确定性计算或非语义类需求（如对账、单据识别、经营报表）被持续追加进 AI 专项，"
        "挤占语义类场景投入并放大资损与审计风险",
        "以 1.5 建设范围边界为唯一裁决依据：此类需求经需求评审会分流至商城系统升级范围，"
        "AI 专项只接规则无法 100% 覆盖的语义与判断类需求"))

    # ── 5 Word 打开时自动更新目录域（新增 1.5 与页码需重算）
    ok &= set_update_fields(doc)

    if not ok:
        print("ABORT: 未保存")
        raise SystemExit(1)
    doc.save(DOC)
    print("saved:", DOC)


if __name__ == "__main__":
    main()
