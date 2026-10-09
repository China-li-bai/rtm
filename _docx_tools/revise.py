"""AI 专项方案 docx 精确修订器。

原则：
  - 只在 run 级别做子串替换，不合并 run、不动样式，避免破坏 Word 格式与目录域。
  - 每条替换必须声明期望命中数，任一不符则整体中止且不落盘（可回滚）。
用法：
  python3 revise.py batch_a        # 文本级逻辑/文案修正
  python3 revise.py verify         # 打印校验摘要
"""
import sys
from docx import Document

DOC = "../AI专项方案v2.docx"


def _runs(doc):
    """正文段落 + 表格单元格内段落的全部 run。"""
    for p in doc.paragraphs:
        for r in p.runs:
            yield p, r
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        yield p, r


def apply(doc, edits):
    """edits: list of (old, new, expect)。old/new 为 run 内子串。"""
    report = []
    for old, new, expect in edits:
        hits = 0
        for p, r in _runs(doc):
            if old in r.text:
                r.text = r.text.replace(old, new)
                hits += 1
        report.append((old[:38], new[:38], expect, hits))
        if hits != expect:
            print(f"  ✗ 命中数不符 期望{expect} 实际{hits}  <- {old[:60]}")
            return report, False
    return report, True


BATCH_A = [
    # ── 文案：图号按出现顺序连续（1.3 协同图在前，2.1 架构图在后）
    ("图 2　三大 AI 中台上下游协同关系", "图 1　三大 AI 中台上下游协同关系", 1),
    ("图 1　AI 专项总体架构", "图 2　AI 专项总体架构", 1),
    # ── 文案：中英混排的节引用符号
    ("详见 §2.6 与 §7.3 的图示说明", "详见 2.6 节与 7.3 节的图示说明", 1),
    # ── 文案：小标题里的实施口语残留（与 6.3.4 体例统一）
    (" （ml+ label studio自动化标注）", "", 1),
    ("（使用ML + label studio 自动化标注）", "", 1),
    # ── 逻辑：AICS-008 被两个功能点重复占用，会话功能实际 5 项（003~007）
    ("（AICS-003 至 AICS-008）", "（AICS-003 至 AICS-007）", 1),
    # ── 逻辑：评测集 ground truth 不能是“自动化标注”（与 8.2 交叉验证原则冲突，会指标虚高）
    #    先处理指标说明整句，再做通用替换，避免子串互相吞匹配
    ("回答符合自动化标注的比例", "回答与人工复核标注答案一致的比例", 1),
    ("自动化标注的", "ML 预标注并经人工交叉复核的", 9),
    # ── 逻辑：与全文云端 API 调用自相矛盾的安全表述
    ("：提供基础大模型 API 能力，不接触业务数据；通过数据处理协议明确数据用途",
     "：提供基础大模型 API 能力，经统一 API 网关调用。供应商不获取业务库全量数据，"
     "仅接收经脱敏与字段最小化处理后的必要内容（商品/价格/订单主数据不出域）；"
     "调用全量留痕，并以数据处理协议明确用途、留存与删除边界", 1),
]


def main():
    cmd = sys.argv[1] if len(sys.argv) > 1 else "verify"
    doc = Document(DOC)
    if cmd == "batch_a":
        report, ok = apply(doc, BATCH_A)
        for old, new, exp, got in report:
            print(f"  {'✓' if exp == got else '✗'} {exp}/{got}  {old!r} -> {new!r}")
        if not ok:
            print("ABORT: 未保存")
            sys.exit(1)
        doc.save(DOC)
        print("saved:", DOC)
    else:
        n_auto = sum(1 for _, r in _runs(doc) if "自动化标注" in r.text)
        n_sec = sum(1 for _, r in _runs(doc) if "§" in r.text)
        print("残留『自动化标注』run 数:", n_auto)
        print("残留『§』run 数:", n_sec)
        for p in doc.paragraphs:
            if p.text.startswith("图 "):
                print("  图题:", p.text[:40])
            if "AICS-00" in p.text and "会话" in p.text:
                print("  客服编号:", p.text[:56])


if __name__ == "__main__":
    main()
